"""把已核验稿件导出为适合发布或复盘的 DOCX / PDF。"""
from __future__ import annotations

from html import escape
from pathlib import Path

from .models import Session
from .norm import usable_images


def _image_map(s: Session) -> dict[str, Path]:
    if not s.evidence:
        return {}
    result: dict[str, Path] = {}
    for image in usable_images(s.evidence, s.brief.product):
        path = Path(image.media_id)
        if path.is_file():
            result[image.id] = path
    return result


def _used_facts(s: Session):
    if not s.draft or not s.evidence:
        return []
    used = {fact_id for paragraph in s.draft.body for fact_id in paragraph.fact_ids}
    return [fact for fact in s.evidence.facts if fact.id in used]


def create_docx(s: Session, target: Path | str, full_report: bool = False) -> Path:
    """生成发布稿或含依据与核验信息的完整 Word 报告。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    if not s.draft:
        raise ValueError("当前会话还没有可导出的稿件")
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = section.right_margin = Cm(2.35)

    styles = doc.styles
    for name, size in (("Normal", 10.5), ("Title", 24), ("Heading 1", 16), ("Heading 2", 13)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].paragraph_format.space_after = Pt(8)
    styles["Normal"].paragraph_format.line_spacing = 1.45
    styles["Heading 1"].paragraph_format.space_before = Pt(16)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)
    styles["Heading 1"].paragraph_format.page_break_before = False
    styles["Heading 1"].paragraph_format.keep_with_next = True
    # Word 默认 Title 样式可能带主题色下边线，导出稿保持干净的黑色标题。
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)

    title = doc.add_paragraph(s.draft.title or f"{s.brief.product}营销文案", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(7)
    meta = doc.add_paragraph(f"{s.brief.platform}  |  {s.brief.audience}  |  {s.brief.scene}")
    meta.runs[0].font.color.rgb = RGBColor(102, 112, 133)
    meta.paragraph_format.space_after = Pt(18)

    if full_report:
        doc.add_heading("一 场景判断", level=1)
        for label, value in (
            ("产品", s.brief.product), ("目标人群", s.brief.audience),
            ("使用场景", s.brief.scene), ("发布平台", s.brief.platform),
            ("内容目标", s.brief.goal), ("核心问题", s.brief.main_problem or "未单独说明"),
        ):
            p = doc.add_paragraph(style="Normal")
            p.add_run(f"{label}：").bold = True
            p.add_run(value)
        doc.add_heading("二 知识库依据", level=1)
        for fact in _used_facts(s):
            p = doc.add_paragraph(style="Normal")
            p.style = doc.styles["Normal"]
            p.add_run(f"{fact.id}  ").bold = True
            p.add_run(f"{fact.text}\n")
            source = p.add_run(f"来源：{fact.source_path or fact.source_title}")
            source.font.size = Pt(8.5)
            source.font.color.rgb = RGBColor(102, 112, 133)
        doc.add_heading("三 图文软文", level=1)

    images = _image_map(s)
    placements: dict[int, list] = {}
    for placement in s.draft.images:
        placements.setdefault(placement.after_paragraph, []).append(placement)

    def add_images(position: int) -> None:
        for placement in placements.get(position, []):
            path = images.get(placement.image_id)
            if not path:
                continue
            try:
                doc.add_picture(str(path), width=Cm(14.8))
                caption = doc.add_paragraph(placement.caption or path.name)
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.runs[0].font.size = Pt(8.5)
                caption.runs[0].font.color.rgb = RGBColor(102, 112, 133)
            except (OSError, ValueError):
                continue

    if s.draft.intro:
        p = doc.add_paragraph(s.draft.intro)
        if p.runs:
            p.runs[0].bold = True
    add_images(0)
    for index, paragraph in enumerate(s.draft.body, 1):
        doc.add_paragraph(paragraph.text)
        add_images(index)
    if s.draft.cta:
        p = doc.add_paragraph(s.draft.cta)
        if p.runs:
            p.runs[0].bold = True
    if s.draft.tags:
        p = doc.add_paragraph(" ".join(f"#{tag.lstrip('#')}" for tag in s.draft.tags))
        p.runs[0].font.color.rgb = RGBColor(64, 104, 180)

    if full_report:
        doc.add_heading("四 核验结果", level=1)
        if s.review and s.review.passed:
            doc.add_paragraph("产品事实、图片型号和明确合规规则均无未解决的阻断项。")
        elif s.review:
            doc.add_paragraph(f"仍有 {len(s.review.errors)} 项阻断问题。")
        if s.review:
            for issue in s.review.issues:
                prefix = "必须处理" if issue.severity == "error" else "提示"
                doc.add_paragraph(f"[{prefix}] {issue.where}：{issue.detail}")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("华为 PC 营销软文智能体")
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(140, 148, 160)
    doc.save(target)
    return target


def create_pdf(s: Session, target: Path | str, full_report: bool = False) -> Path:
    """生成发布稿或含依据与核验信息的完整 PDF。"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer

    if not s.draft:
        raise ValueError("当前会话还没有可导出的稿件")
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    base = getSampleStyleSheet()
    body = ParagraphStyle(
        "ChineseBody", parent=base["BodyText"], fontName="STSong-Light", fontSize=10.5,
        leading=17, textColor=colors.HexColor("#182230"), spaceAfter=8,
    )
    title_style = ParagraphStyle(
        "ChineseTitle", parent=body, fontSize=22, leading=29, textColor=colors.black, spaceAfter=8,
    )
    heading = ParagraphStyle(
        "ChineseHeading", parent=body, fontSize=15, leading=21, textColor=colors.black,
        spaceBefore=14, spaceAfter=8,
    )
    muted = ParagraphStyle(
        "ChineseMuted", parent=body, fontSize=8.5, leading=13, textColor=colors.HexColor("#667085"),
    )
    tag_style = ParagraphStyle(
        "ChineseTags", parent=body, textColor=colors.HexColor("#4068B4"), spaceBefore=4,
    )

    def page_footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("STSong-Light", 8)
        canvas.setFillColor(colors.HexColor("#8C94A0"))
        canvas.drawCentredString(A4[0] / 2, 12 * mm, f"华为 PC 营销软文智能体  ·  {doc.page}")
        canvas.restoreState()

    story = [
        Paragraph(escape(s.draft.title or f"{s.brief.product}营销文案"), title_style),
        Paragraph(escape(f"{s.brief.platform}  |  {s.brief.audience}  |  {s.brief.scene}"), muted),
        Spacer(1, 6 * mm),
    ]
    if full_report:
        story.append(Paragraph("一 场景判断", heading))
        for label, value in (
            ("产品", s.brief.product), ("目标人群", s.brief.audience),
            ("使用场景", s.brief.scene), ("发布平台", s.brief.platform),
            ("内容目标", s.brief.goal), ("核心问题", s.brief.main_problem or "未单独说明"),
        ):
            story.append(Paragraph(f"<b>{escape(label)}：</b>{escape(value)}", body))
        story.append(Paragraph("二 知识库依据", heading))
        for fact in _used_facts(s):
            story.append(Paragraph(f"<b>{escape(fact.id)}</b>  {escape(fact.text)}", body))
            story.append(Paragraph(f"来源：{escape(fact.source_path or fact.source_title)}", muted))
        story.append(Paragraph("三 图文软文", heading))

    images = _image_map(s)
    placements: dict[int, list] = {}
    for placement in s.draft.images:
        placements.setdefault(placement.after_paragraph, []).append(placement)

    def append_images(position: int) -> None:
        for placement in placements.get(position, []):
            path = images.get(placement.image_id)
            if not path:
                continue
            try:
                image = Image(str(path))
                image._restrictSize(155 * mm, 92 * mm)
                image.hAlign = "CENTER"
                story.extend([image, Paragraph(escape(placement.caption or path.name), muted), Spacer(1, 3 * mm)])
            except (OSError, ValueError):
                continue

    if s.draft.intro:
        story.append(Paragraph(f"<b>{escape(s.draft.intro)}</b>", body))
    append_images(0)
    for index, paragraph in enumerate(s.draft.body, 1):
        story.append(Paragraph(escape(paragraph.text).replace("\n", "<br/>"), body))
        append_images(index)
    if s.draft.cta:
        story.append(Paragraph(f"<b>{escape(s.draft.cta)}</b>", body))
    if s.draft.tags:
        story.append(Paragraph(escape(" ".join(f"#{tag.lstrip('#')}" for tag in s.draft.tags)), tag_style))

    if full_report:
        story.append(Paragraph("四 核验结果", heading))
        if s.review and s.review.passed:
            story.append(Paragraph("产品事实、图片型号和明确合规规则均无未解决的阻断项。", body))
        elif s.review:
            story.append(Paragraph(f"仍有 {len(s.review.errors)} 项阻断问题。", body))
        if s.review:
            for issue in s.review.issues:
                prefix = "必须处理" if issue.severity == "error" else "提示"
                story.append(Paragraph(escape(f"[{prefix}] {issue.where}：{issue.detail}"), body))

    document = SimpleDocTemplate(
        str(target), pagesize=A4, rightMargin=22 * mm, leftMargin=22 * mm,
        topMargin=20 * mm, bottomMargin=21 * mm, title=s.draft.title,
    )
    document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return target
